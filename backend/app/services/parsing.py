"""Text extraction for knowledge-base documents.

Each parser returns sections: (text, page) pairs, where page is 1-based for PDFs and
None otherwise, so search results can say where a passage came from.

A PDF with no text layer (a scan) is read with Tesseract OCR when it's installed (the LLMCoach
image includes it). OCR takes seconds a page, so its result is cached by the file's hash: the
inbox's check and the indexing job read a scan once between them.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".md", ".markdown", ".txt", ".text", ".rst", ".csv", ".json", ".html", ".htm", ".docx"}

Section = tuple[str, int | None]


class ParseError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    # UTF-16 only with a BOM: almost any even-length byte string "decodes" as UTF-16, which
    # would turn Windows-1252 / Latin-1 text into CJK garbage.
    encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp1252")
    for enc in encodings:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def ocr_available() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def _ocr_pdf(path: Path) -> list[Section]:
    """Reads the page images of a scanned PDF with Tesseract."""
    import pytesseract
    from pypdf import PdfReader

    out: list[Section] = []
    for i, page in enumerate(PdfReader(str(path)).pages):
        texts = []
        for img in page.images:
            try:
                texts.append(pytesseract.image_to_string(img.image))
            except Exception as e:  # an image Pillow or Tesseract can't read: skip it
                log.warning("OCR skipped an image on page %s of %s: %s", i + 1, path.name, e)
        out.append(("\n".join(t.strip() for t in texts if t.strip()), i + 1))
    return out


def _ocr_cached(path: Path) -> list[Section]:
    from ..config import settings

    cache = settings.data_dir / "ocr" / f"{hashlib.sha256(path.read_bytes()).hexdigest()}.json"
    try:
        return [(t, p) for t, p in json.loads(cache.read_text(encoding="utf-8"))]
    except (OSError, ValueError):
        pass
    sections = _ocr_pdf(path)
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(sections), encoding="utf-8")
    return sections


def parse(path: Path, ocr: bool = True) -> list[Section]:
    """`ocr=False` skips reading scans, for callers someone is waiting on (an upload's check)."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ParseError(f"unsupported file type {ext or '(none)'}; supported: {', '.join(sorted(SUPPORTED))}")
    if ext == ".pdf":
        sections = _pdf(path)
    elif ext == ".docx":
        sections = [(_docx(path), None)]
    elif ext in (".html", ".htm"):
        sections = [(_html(_read_text(path)), None)]
    else:
        sections = [(_read_text(path), None)]
    sections = [(_clean(t), p) for t, p in sections]
    sections = [(t, p) for t, p in sections if t]
    tried_ocr = False
    if not sections and ext == ".pdf" and ocr and ocr_available():
        tried_ocr = True
        sections = [(t, p) for t, p in ((_clean(t), p) for t, p in _ocr_cached(path)) if t]
    if not sections:
        hint = ""
        if ext == ".pdf":
            hint = (" Even OCR found no text in it." if tried_ocr else
                    " It looks like a scan with no text layer; reading scans needs Tesseract OCR, which the LLMCoach image includes."
                    if ocr else " It looks like a scan; it's read with OCR when it's indexed.")
        raise ParseError(f"no text found in {path.name}.{hint}")
    return sections


def _pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            reader.decrypt("")
        return [(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]
    except Exception as e:  # pypdf raises many types for damaged files
        raise ParseError(f"couldn't read PDF: {e}") from e


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(parts)


def _html(html: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")
