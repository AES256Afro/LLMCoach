import io
import time

import docx
import pytest

from app.services.chunking import chunk_sections
from app.services.parsing import ParseError, parse
from tests.test_jobs import wait_final


def minimal_pdf(pages: list[str]) -> bytes:
    """A valid, uncompressed PDF with one line of Helvetica text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out.encode()))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out.encode())
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode()


# ---- chunking ---------------------------------------------------------------------

def test_chunks_respect_size_and_overlap():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(20))
    chunks = chunk_sections([(text, None)], size=500, overlap=80)
    assert len(chunks) > 5
    assert all(len(c.text) <= 500 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    # Consecutive chunks share text (the overlap).
    assert any(chunks[i].text[-30:].split()[-1] in chunks[i + 1].text[:120] for i in range(len(chunks) - 1))


def test_chunks_keep_page_numbers():
    chunks = chunk_sections([("alpha " * 50, 1), ("beta " * 50, 2)], size=200, overlap=0)
    assert {c.page for c in chunks} == {1, 2}
    assert all(("alpha" in c.text) == (c.page == 1) for c in chunks)


def test_unbreakable_text_is_hard_split():
    chunks = chunk_sections([("x" * 2500, None)], size=1000, overlap=0)
    assert all(len(c.text) <= 1000 for c in chunks) and sum(len(c.text) for c in chunks) == 2500


# ---- parsing ----------------------------------------------------------------------

def test_parse_formats(tmp_path):
    (tmp_path / "a.md").write_text("# Title\n\nSome **markdown**.", encoding="utf-8")
    (tmp_path / "b.html").write_text("<html><script>x()</script><body><h1>Hi</h1><p>There</p></body></html>")
    d = docx.Document()
    d.add_paragraph("Docx paragraph")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "cell A", "cell B"
    d.save(tmp_path / "c.docx")
    (tmp_path / "d.pdf").write_bytes(minimal_pdf(["First page text", "Second page text"]))

    assert "Some **markdown**" in parse(tmp_path / "a.md")[0][0]
    html = parse(tmp_path / "b.html")[0][0]
    assert "Hi" in html and "There" in html and "x()" not in html
    assert "cell A | cell B" in parse(tmp_path / "c.docx")[0][0]
    pdf = parse(tmp_path / "d.pdf")
    assert [p for _, p in pdf] == [1, 2] and "Second page" in pdf[1][0]


def test_parse_text_encodings(tmp_path):
    text = "Café crème — naïve “quotes” at 20°C. "
    cp = (text * 4).encode("cp1252")
    assert len(cp) % 2 == 0  # even length: the case utf-16 would happily (and wrongly) decode
    (tmp_path / "cp.txt").write_bytes(cp)
    (tmp_path / "u16.txt").write_bytes((text * 4).encode("utf-16"))  # with BOM
    (tmp_path / "u8.txt").write_bytes((text * 4).encode("utf-8"))
    for name in ("cp.txt", "u16.txt", "u8.txt"):
        assert parse(tmp_path / name)[0][0].startswith("Café crème — naïve “quotes” at 20°C."), name


def test_parse_rejects_unsupported_and_empty(tmp_path):
    (tmp_path / "x.exe").write_bytes(b"MZ")
    (tmp_path / "e.txt").write_text("   \n  ")
    with pytest.raises(ParseError, match="unsupported"):
        parse(tmp_path / "x.exe")
    with pytest.raises(ParseError, match="no text"):
        parse(tmp_path / "e.txt")


# ---- end to end: upload -> ingest job -> search -------------------------------------

def test_upload_ingest_search_delete(client, fake):
    pid = client.post("/api/projects", json={"name": "kb-e2e"}).json()["id"]
    files = [
        ("files", ("zoo.txt", io.BytesIO(b"zebras zigzag in the zoo. " * 30), "text/plain")),
        ("files", ("fruit.md", io.BytesIO(b"# Fruit\n\napples and bananas are a healthy snack. " * 30), "text/markdown")),
        ("files", ("dupe.txt", io.BytesIO(b"zebras zigzag in the zoo. " * 30), "text/plain")),  # same content
        ("files", ("tool.exe", io.BytesIO(b"MZ"), "application/octet-stream")),
    ]
    r = client.post(f"/api/projects/{pid}/documents", files=files)
    assert r.status_code == 201, r.text
    out = r.json()
    assert [d["filename"] for d in out["documents"]] == ["zoo.txt", "fruit.md"]
    assert {s["reason"] for s in out["skipped"]} == {"already in this knowledge base", "unsupported file type"}

    job = wait_final(client, out["job"]["id"])
    assert job["status"] == "done", client.get(f"/api/jobs/{job['id']}/log").text
    docs = {d["filename"]: d for d in client.get(f"/api/projects/{pid}/documents").json()}
    assert all(d["status"] == "ready" and d["chunk_count"] > 0 for d in docs.values())
    assert docs["zoo.txt"]["embed_model"] == "ollama/nomic-embed-text"
    # nomic's document prefix was applied when embedding.
    assert any(t.startswith("search_document: zebras") for t in fake.app.state.embed_inputs)

    stats = client.get(f"/api/projects/{pid}/knowledge").json()
    assert stats["documents"] == 2 and stats["chunks"] > 2 and stats["dimension"] == 8 and not stats["stale_doc_ids"]

    res = client.post(f"/api/projects/{pid}/search", json={"query": "zebras zigzag zoo", "top_k": 3}).json()
    hits = res["results"]
    assert res["mode"] == "hybrid"
    assert hits[0]["filename"] == "zoo.txt" and 0 < hits[0]["score"] <= 1
    client.patch(f"/api/projects/{pid}", json={"settings": {"search_mode": "vector"}})
    res = client.post(f"/api/projects/{pid}/search", json={"query": "zebras zigzag zoo", "top_k": 3}).json()
    assert res["mode"] == "vector" and res["results"][0]["filename"] == "zoo.txt"
    client.patch(f"/api/projects/{pid}", json={"settings": {"search_mode": None}})
    only_fruit = client.post(f"/api/projects/{pid}/search",
                             json={"query": "zebras", "doc_ids": [docs["fruit.md"]["id"]]}).json()["results"]
    assert {h["filename"] for h in only_fruit} == {"fruit.md"}

    chunks = client.get(f"/api/projects/{pid}/documents/{docs['zoo.txt']['id']}/chunks").json()
    assert chunks["total"] == docs["zoo.txt"]["chunk_count"] and chunks["chunks"][0]["chunk_index"] == 0

    # Changing the embedding model marks documents stale.
    client.patch(f"/api/projects/{pid}", json={"settings": {"embed_model": "ollama/other-embed"}})
    assert len(client.get(f"/api/projects/{pid}/knowledge").json()["stale_doc_ids"]) == 2
    client.patch(f"/api/projects/{pid}", json={"settings": {"embed_model": None}})

    assert client.delete(f"/api/projects/{pid}/documents/{docs['zoo.txt']['id']}").status_code == 204
    assert client.get(f"/api/projects/{pid}/knowledge").json()["documents"] == 1
    assert client.delete(f"/api/projects/{pid}").status_code == 204


def test_failed_document_does_not_block_others(client):
    pid = client.post("/api/projects", json={"name": "kb-partial"}).json()["id"]
    files = [("files", ("empty.txt", io.BytesIO(b"   \n\n   "), "text/plain")),
             ("files", ("good.txt", io.BytesIO(b"useful text " * 50), "text/plain"))]
    out = client.post(f"/api/projects/{pid}/documents", files=files).json()
    assert wait_final(client, out["job"]["id"])["status"] == "done"
    docs = {d["filename"]: d for d in client.get(f"/api/projects/{pid}/documents").json()}
    assert docs["good.txt"]["status"] == "ready"
    assert docs["empty.txt"]["status"] == "failed" and "no text" in docs["empty.txt"]["error"]


def test_cancelled_ingest_marks_documents_failed(client):
    pid = client.post("/api/projects", json={"name": "kb-cancel"}).json()["id"]
    # Occupy the single job slot so the ingest job stays queued, then cancel it.
    blocker = client.post("/api/jobs", json={"kind": "demo", "config": {"steps": 1000, "delay": 0.05}}).json()
    out = client.post(f"/api/projects/{pid}/documents",
                      files=[("files", ("a.txt", io.BytesIO(b"text " * 100), "text/plain"))]).json()
    # A document waiting for a queued ingest job can't be deleted from under it.
    r = client.delete(f"/api/projects/{pid}/documents/{out['documents'][0]['id']}")
    assert r.status_code == 409 and "queued" in r.json()["detail"]
    client.post(f"/api/jobs/{out['job']['id']}/cancel")
    client.post(f"/api/jobs/{blocker['id']}/cancel")
    wait_final(client, blocker["id"])
    deadline = time.time() + 5
    while time.time() < deadline:
        doc = client.get(f"/api/projects/{pid}/documents").json()[0]
        if doc["status"] == "failed":
            break
        time.sleep(0.1)
    assert doc["status"] == "failed" and "cancelled" in doc["error"]


def test_markdown_sections_start_new_chunks():
    md = "# Intro\n\n" + "intro text. " * 30 + "\n\n## Backups\n\n" + "backup text. " * 30 + "\n\n## Network\n\n" + "net text. " * 30
    chunks = chunk_sections([(md, None)], size=600, overlap=0)
    assert any(c.text.startswith("## Backups") for c in chunks)
    assert not any("intro text" in c.text and "backup text" in c.text for c in chunks)
